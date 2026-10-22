"""
올리브영 메이크업(립/베이스/아이) 옵션별 대표 색상 수집 PoC  (v5)
- 비로그인, API key 없음, requests 계열(curl_cffi) + BeautifulSoup, Selenium/Playwright 없음
- 중분류당 상품 10개, oliveyoung.co.kr 요청 간격 5초
- 상세 페이지 HTML 안의 Next.js 데이터(self.__next_f.push)에서 옵션/브랜드/소분류를 읽음
  (화면의 옵션 목록은 JS가 그리기 때문에 HTML 태그 셀렉터로는 비어 보임)
- 컬러칩 이미지는 메모리에서만 열어 색을 뽑고 저장하지 않음
- 결과: 15열 엑셀, '색상(배경)' 칸을 HEX 색으로 채움

pip install curl_cffi beautifulsoup4 pillow openpyxl
같은 폴더에 cookie.txt (cookie / user-agent 형식) 필요
"""
import io, os, re, json, time, random
from datetime import datetime

from bs4 import BeautifulSoup
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

try:
    from curl_cffi import requests as http
    IMPERSONATE = {"impersonate": "chrome"}
except ImportError:
    import requests as http
    IMPERSONATE = {}

BASE = "https://www.oliveyoung.co.kr"
MAKEUP_URL = f"{BASE}/store/display/getCategoryShop.do?dispCatNo=10000010002"
TARGETS = {"립메이크업": "100000100020006",
           "베이스메이크업": "100000100020001",
           "아이메이크업": "100000100020007"}
PER_CATEGORY = 10
PAGE_DELAY = 5            # 올리브영 페이지 요청 간격(초)
IMG_DELAY = 1             # 컬러칩 이미지(CDN) 요청 간격(초)
IMG_SIZE = 300            # 분석용 컬러칩 크기
MIN_COLOR_FRAC = 0.5      # 유채색 비율이 이보다 낮으면 제품 사진으로 보고 HEX 제외 (테스트 스크립트와 동일)
HERE = os.path.dirname(os.path.abspath(__file__))
OUT_XLSX = os.path.join(HERE, f"oliveyoung_makeup_colors_{datetime.now():%Y%m%d_%H%M}.xlsx")

COLUMNS = ["상품번호", "중분류", "제품유형(소분류)", "브랜드명", "상품명", "호수", "호수이름",
           "올리브영 옵션명", "올리브영 HEX", "색상(배경)", "L*", "a*", "b*",
           "옵션 이미지 URL", "수집일시"]


# ───────── 세션 (cookie.txt) ─────────
def load_cookie_file():
    for name in ("cookie.txt", "cookie.txt.txt", "cookie"):
        p = os.path.join(HERE, name)
        if os.path.isfile(p):
            raw = open(p, "rb").read()
            for enc in ("utf-8-sig", "utf-16", "cp949"):
                try:
                    text = raw.decode(enc); break
                except UnicodeError:
                    continue
            lines = [l.strip() for l in text.splitlines() if l.strip()]
            vals = {lines[i].lower(): lines[i + 1] for i in range(len(lines) - 1)
                    if lines[i].lower() in ("cookie", "user-agent")}
            jar = {}
            for part in vals.get("cookie", "").split(";"):
                if "=" in part:
                    k, v = part.strip().split("=", 1)
                    jar.pop(k, None); jar[k] = v
            print(f"[쿠키] {name} 읽음 → {len(jar)}개, cf_clearance {'있음' if 'cf_clearance' in jar else '없음'}")
            return "; ".join(f"{k}={v}" for k, v in jar.items()), vals.get("user-agent", "")
    print(f"[쿠키] cookie.txt 없음 ({HERE})")
    return "", ""


COOKIE, UA = load_cookie_file()
session = http.Session()
session.headers.update({"Accept-Language": "ko-KR,ko;q=0.9,en-US;q=0.8",
                        "Referer": f"{BASE}/store/main/main.do"})
if UA:
    session.headers["User-Agent"] = UA
if COOKIE:
    session.headers["Cookie"] = COOKIE
_last = {"page": 0.0, "img": 0.0}


def get(url, kind="page", **kw):
    wait = (PAGE_DELAY if kind == "page" else IMG_DELAY) - (time.time() - _last[kind])
    if wait > 0:
        time.sleep(wait + random.uniform(0, 0.5))
    r = session.get(url, timeout=20, **IMPERSONATE, **kw)
    _last[kind] = time.time()
    if r.status_code == 403:
        raise RuntimeError(f"차단됨(403, {r.headers.get('server', '?')}) → cookie.txt 갱신 필요")
    r.raise_for_status()
    return r


# ───────── 1. 목록 페이지 → 상품번호 ─────────
def list_goods(cat_no, limit):
    params = {"dispCatNo": cat_no, "fltDispCatNo": "", "prdSort": "01", "pageIdx": 1, "rowsPerPage": 24}
    soup = BeautifulSoup(get(f"{BASE}/store/display/getMCategoryList.do", params=params,
                             headers={"Referer": MAKEUP_URL}).text, "html.parser")
    goods = []
    for a in soup.select("a[data-ref-goodsno], a[href*='goodsNo=']"):
        no = a.get("data-ref-goodsno") or (re.search(r"goodsNo=(\w+)", a.get("href", "")) or [None, None])[1]
        if no and no not in goods:
            goods.append(no)
        if len(goods) >= limit:
            break
    return goods


# ───────── 2. 상세 페이지 → Next.js 데이터 파싱 ─────────
PUSH_RE = re.compile(r'self\.__next_f\.push\(\[1,("(?:[^"\\]|\\.)*")\]\)')


def next_data_text(html):
    """self.__next_f.push([1,"..."]) 조각들을 이어 붙여 하나의 텍스트로 복원"""
    return "".join(json.loads(m.group(1)) for m in PUSH_RE.finditer(html))


def json_after(txt, key):
    """txt 안에서 "key": 바로 뒤의 JSON 값(배열/객체)을 파싱"""
    i = txt.find(f'"{key}":')
    if i < 0:
        return None
    try:
        return json.JSONDecoder().raw_decode(txt, i + len(key) + 3)[0]
    except ValueError:
        return None


def str_after(txt, key):
    m = re.search(rf'"{key}":"((?:[^"\\]|\\.)*)"', txt)
    return json.loads(f'"{m.group(1)}"') if m else ""


def parse_detail(html, goods_no):
    txt = next_data_text(html)
    cat = json_after(txt, "displayCategory") or {}
    options = [o for o in (json_after(txt, "options") or []) if o.get("goodsNumber") == goods_no]
    seen, uniq = set(), []
    for o in options:                                  # 같은 옵션이 두 번 나오는 경우 제거
        if o.get("optionNumber") not in seen:
            seen.add(o.get("optionNumber")); uniq.append(o)
    name = (uniq[0].get("goodsName") if uniq else "") or str_after(txt, "goodsName")
    return {"sub_cat": cat.get("leafCategoryName") or cat.get("lowerCategoryName", ""),
            "brand": str_after(txt, "onlineBrandName"),
            "name": name, "options": uniq}


# ───────── 3. 호수 분리 / 색 추출 / Lab ─────────
HOSU_RE = re.compile(r"^#?([A-Za-z]{0,3}\d+[A-Za-z0-9.\-]*)[\s._)\]:-]*(.*)$")


def split_hosu(option_name):
    s = re.sub(r"\[[^\]]*\]|\([^)]*\)", " ", option_name).strip()   # [기획], (단품) 태그 제거
    m = HOSU_RE.match(s)
    return (m.group(1), m.group(2).strip()) if m else ("", s)


BG_DIST = 40          # 배경색과 이 거리 이내면 배경으로 간주
MIN_OBJECT_FRAC = 0.02  # 배경을 뺀 물체가 이미지의 2% 미만이면 색 없음


def extract_color(img_bytes):
    """
    1) 테두리 픽셀의 중앙값으로 배경색 추정 → 배경 픽셀 제거 (흰/회색 배경 위 작은 스와치 대응)
    2) 남은 '물체' 픽셀 중 유채색 비율이 MIN_COLOR_FRAC 미만이면 제품 사진으로 보고 제외
    3) 유채색 픽셀을 밝기순 정렬 → 30~70% 구간 평균 (하이라이트/그림자/펄 제외)
    """
    im = Image.open(io.BytesIO(img_bytes))
    if im.mode in ("RGBA", "LA", "P"):                       # 투명 배경은 흰색으로 합성
        im = im.convert("RGBA")
        bg_img = Image.new("RGBA", im.size, (255, 255, 255, 255))
        im = Image.alpha_composite(bg_img, im)
    im = im.convert("RGB")
    im.thumbnail((200, 200))
    w, h = im.size
    pix = im.load()

    border = [pix[x, 0] for x in range(w)] + [pix[x, h - 1] for x in range(w)] + \
             [pix[0, y] for y in range(h)] + [pix[w - 1, y] for y in range(h)]
    bg = [sorted(p[k] for p in border)[len(border) // 2] for k in range(3)]

    bg_is_color = max(bg) - min(bg) >= 25        # 테두리가 유채색 = 이미지 전체가 컬러칩
    total, obj = 0, []
    for y in range(0, h, 2):
        for x in range(0, w, 2):
            total += 1
            r, g, b = pix[x, y]
            if bg_is_color or ((r - bg[0]) ** 2 + (g - bg[1]) ** 2 + (b - bg[2]) ** 2) ** 0.5 > BG_DIST:
                obj.append((r, g, b))
    if len(obj) < max(30, total * MIN_OBJECT_FRAC):
        return None

    col = [(r, g, b, 0.2126 * r + 0.7152 * g + 0.0722 * b) for r, g, b in obj if max(r, g, b) - min(r, g, b) >= 25]
    if len(col) / len(obj) < MIN_COLOR_FRAC or len(col) < 30:
        return None
    col.sort(key=lambda p: p[3])
    sel = col[int(len(col) * 0.3):int(len(col) * 0.7)] or col
    return [sum(p[j] for p in sel) / len(sel) for j in range(3)]


def to_lab(rgb):
    def lin(v):
        v /= 255
        return v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4
    R, G, B = (lin(c) for c in rgb)
    X = (R * 0.4124 + G * 0.3576 + B * 0.1805) / 0.95047
    Y = R * 0.2126 + G * 0.7152 + B * 0.0722
    Z = (R * 0.0193 + G * 0.1192 + B * 0.9505) / 1.08883
    f = lambda t: t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116
    return round(116 * f(Y) - 16, 1), round(500 * (f(X) - f(Y)), 1), round(200 * (f(Y) - f(Z)), 1)


def image_url(o):
    img = o.get("colorChipImage") or o.get("optionImage") or {}
    return (img.get("url", "") + img.get("path", "")) if img.get("path") else ""


# ───────── 4. 상품 1개 → 행 목록 ─────────
def crawl_goods(mid_name, goods_no):
    d = parse_detail(get(f"{BASE}/store/goods/getGoodsDetail.do?goodsNo={goods_no}").text, goods_no)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    base = [goods_no, mid_name, d["sub_cat"], d["brand"], d["name"]]
    if not d["options"]:
        return [base + ["", "", "단일상품(옵션없음)", "", "", "", "", "", "", now]]
    rows = []
    n_opt = len(d["options"])
    for k, o in enumerate(d["options"], 1):
        print(f"\r      컬러칩 {k}/{n_opt} 분석 중...", end="", flush=True)
        opt = o.get("optionName", "")
        hosu, hosu_name = split_hosu(opt)
        url = image_url(o)
        hexv, lab = "", ("", "", "")
        if url:
            try:
                rgb = extract_color(get(f"{url}?RS={IMG_SIZE}x0&QT=100", kind="img").content)
                if rgb:
                    hexv = "#%02X%02X%02X" % tuple(round(c) for c in rgb)
                    lab = to_lab(rgb)
            except Exception as e:
                print(f"      이미지 실패({opt}): {e}")
        rows.append(base + [hosu, hosu_name, opt, hexv, "", *lab, url, now])
    print("\r" + " " * 40 + "\r", end="")
    return rows


def fmt_dur(sec):
    sec = int(max(sec, 0))
    return f"{sec // 3600}시간 {sec % 3600 // 60}분" if sec >= 3600 else f"{sec // 60}분 {sec % 60}초"


def eta_line(done, total, started, durations, rough_per_goods):
    """지금까지 상품 1개당 실제 걸린 시간 평균 × 남은 상품 수 (처음엔 대략값)"""
    per = sum(durations) / len(durations) if durations else rough_per_goods
    remain = per * (total - done)
    end = datetime.fromtimestamp(time.time() + remain)
    return (f"진행 {done}/{total} | 경과 {fmt_dur(time.time() - started)} | "
            f"남은 시간 약 {fmt_dur(remain)} (예상 완료 {end:%H:%M})")


def save_excel(rows):
    wb = Workbook(); ws = wb.active; ws.title = "oliveyoung_colors"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    hex_i, fill_i = COLUMNS.index("올리브영 HEX"), COLUMNS.index("색상(배경)") + 1
    for r in rows:
        ws.append(r)
        if r[hex_i]:
            ws.cell(ws.max_row, fill_i).fill = PatternFill("solid", fgColor=r[hex_i][1:])
    for i, w in enumerate([15, 13, 14, 12, 45, 7, 18, 32, 12, 10, 7, 7, 7, 70, 19], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "B2"
    wb.save(OUT_XLSX)
    print(f"\n저장 완료: {OUT_XLSX} ({len(rows)}행, HEX 추출 {sum(1 for r in rows if r[hex_i])}개)")


def main():
    print("HTTP:", "curl_cffi(chrome)" if IMPERSONATE else "requests", "| 쿠키:", bool(COOKIE))
    started = time.time()

    # 1단계: 세 카테고리 목록을 먼저 받아 전체 작업량 확정
    jobs = []
    for mid_name, cat_no in TARGETS.items():
        try:
            goods = list_goods(cat_no, PER_CATEGORY)
            jobs += [(mid_name, no) for no in goods]
            print(f"[{mid_name}] 상품 {len(goods)}개 확인")
        except Exception as e:
            print(f"[{mid_name}] 목록 실패: {e}")

    # 대략값: 상세 페이지 5초 + 옵션 약 20개 × 이미지 1초+α  (첫 상품 이후엔 실측 평균으로 교체)
    rough = PAGE_DELAY + 20 * (IMG_DELAY + 0.3)
    print(f"\n총 {len(jobs)}개 상품 | 첫 예상: 약 {fmt_dur(rough * len(jobs))} "
          f"(옵션 수에 따라 달라지며 진행하면서 보정됩니다)")

    # 2단계: 상품별 수집
    rows, durations, cur_mid = [], [], None
    for i, (mid_name, no) in enumerate(jobs, 1):
        if mid_name != cur_mid:
            print(f"\n[{mid_name}]"); cur_mid = mid_name
        t0 = time.time()
        try:
            got = crawl_goods(mid_name, no)
            rows += got
            print(f"   {i:02d}. {no} [{got[0][2]}] {got[0][3]} {got[0][4][:28]} → 옵션 {len(got)}개, "
                  f"HEX {sum(1 for r in got if r[8])}개")
        except Exception as e:
            print(f"   {i:02d}. {no} 실패: {e}")
        durations.append(time.time() - t0)
        print("       " + eta_line(i, len(jobs), started, durations, rough))
    print(f"\n총 소요 시간: {fmt_dur(time.time() - started)}")
    if rows:
        save_excel(rows)
    else:
        print("수집된 데이터 없음")


if __name__ == "__main__":
    main()
