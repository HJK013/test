"""
수집 스크립트(v8 엑셀) vs Aside 결과(CSV) 색상 비교
사용법:
    python compare_aside_vs_script.py <v8결과.xlsx> <aside_result.csv>
결과 엑셀 시트 2개:
- 'Aside 결과(v8 형식)': Aside CSV를 v8과 같은 19열로 변환 (L*a*b* 계산, 색상(배경) 칠하기)
- '비교': '상품번호 + 올리브영 옵션명'으로 짝지어 두 HEX의 색 차이 ΔE(CIE76)와 두 색을 나란히 표시
  ΔE 기준(대략): 0~2 거의 같음 / 2~5 자세히 보면 다름 / 5~10 다름 / 10+ 확실히 다름
pip install openpyxl
"""
V8_COLUMNS = ["상품번호", "중분류", "제품유형(소분류)", "브랜드명", "상품명", "호수", "호수이름",
              "올리브영 옵션명", "올리브영 HEX", "색상(배경)", "L*", "a*", "b*",
              "옵션 이미지 URL", "수집일시", "제형", "판정 근거", "색 출처", "색 신뢰도"]
import csv, os, re, sys
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


def norm(s):
    return re.sub(r"\s+", "", str(s or "")).lower()


def hex_ok(h):
    return bool(re.fullmatch(r"#?[0-9A-Fa-f]{6}", str(h or "").strip()))


def to_lab(h):
    h = h.strip().lstrip("#")
    rgb = [int(h[i:i + 2], 16) for i in (0, 2, 4)]
    def lin(v):
        v /= 255
        return v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4
    R, G, B = (lin(c) for c in rgb)
    X = (R * 0.4124 + G * 0.3576 + B * 0.1805) / 0.95047
    Y = R * 0.2126 + G * 0.7152 + B * 0.0722
    Z = (R * 0.0193 + G * 0.1192 + B * 0.9505) / 1.08883
    f = lambda t: t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116
    return 116 * f(Y) - 16, 500 * (f(X) - f(Y)), 200 * (f(Y) - f(Z))


def delta_e(h1, h2):
    a, b = to_lab(h1), to_lab(h2)
    return sum((a[i] - b[i]) ** 2 for i in range(3)) ** 0.5


def verdict(de):
    return "거의 같음" if de < 2 else "약간 다름" if de < 5 else "다름" if de < 10 else "확실히 다름"


def load_script(path):
    ws = load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    head = [str(h) for h in rows[0]]
    idx = {h: i for i, h in enumerate(head)}
    out = {}
    for r in rows[1:]:
        key = (norm(r[idx["상품번호"]]), norm(r[idx["올리브영 옵션명"]]))
        out[key] = {"상품명": r[idx["상품명"]], "옵션명": r[idx["올리브영 옵션명"]],
                    "HEX": r[idx["올리브영 HEX"]] or "",
                    "출처": r[idx["색 출처"]] if "색 출처" in idx else "",
                    "제형": r[idx["제형"]] if "제형" in idx else ""}
    return out


def read_aside_rows(path):
    for enc in ("utf-8-sig", "cp949"):
        try:
            with open(path, encoding=enc, newline="") as f:
                return [{(k or "").strip(): (v or "").strip() for k, v in r.items()} for r in csv.DictReader(f)]
        except UnicodeError:
            continue
    return []


def load_aside(rows):
    out = {}
    for r in rows:
        key = (norm(r.get("상품번호")), norm(r.get("올리브영 옵션명")))
        out[key] = {"HEX": r.get("올리브영 HEX") or r.get("HEX") or "",
                    "출처": r.get("색 출처", ""),
                    "측정방법": r.get("측정 방법", ""), "제형": r.get("제형", "")}
    return out


def write_aside_v8(wb, rows):
    """Aside CSV → v8과 같은 19열 시트 (+ 측정 방법). L*a*b*는 HEX로 계산, 색상(배경) 칠함"""
    ws = wb.create_sheet("Aside 결과(v8 형식)")
    cols = V8_COLUMNS + ["측정 방법"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        h = r.get("올리브영 HEX") or r.get("HEX") or ""
        if hex_ok(h):
            h = "#" + h.lstrip("#").upper()
            L, A, B = (round(v, 1) for v in to_lab(h))
        else:
            L = A = B = ""
        vals = {**r, "올리브영 HEX": h, "색상(배경)": "", "L*": L, "a*": A, "b*": B}
        ws.append([vals.get(c, "") for c in cols])
        if hex_ok(h):
            ws.cell(ws.max_row, cols.index("색상(배경)") + 1).fill = PatternFill("solid", fgColor=h.lstrip("#"))
    for i, w in enumerate([15, 13, 14, 12, 45, 7, 18, 32, 12, 10, 7, 7, 7, 70, 19, 10, 45, 50, 12, 14], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "B2"


def main():
    if len(sys.argv) < 3:
        print(__doc__); return
    aside_rows = read_aside_rows(sys.argv[2])
    s, a = load_script(sys.argv[1]), load_aside(aside_rows)
    keys = [k for k in a if k in s] + [k for k in a if k not in s]
    wb = Workbook(); ws = wb.active; ws.title = "비교"
    head = ["상품번호", "상품명", "옵션명", "스크립트 HEX", "스크립트 색", "Aside HEX", "Aside 색",
            "ΔE", "판정", "스크립트 출처", "Aside 측정 방법", "스크립트 제형", "Aside 제형"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    stats = {}
    for k in keys:
        sv, av = s.get(k, {}), a[k]
        de = round(delta_e(sv["HEX"], av["HEX"]), 1) if sv and hex_ok(sv.get("HEX")) and hex_ok(av["HEX"]) else ""
        v = verdict(de) if de != "" else ("스크립트에 없음" if not sv else "한쪽 색 없음")
        stats[v] = stats.get(v, 0) + 1
        ws.append([k[0].upper(), sv.get("상품명", ""), sv.get("옵션명", k[1]), sv.get("HEX", ""), "",
                   av["HEX"], "", de, v, sv.get("출처", ""), av["측정방법"], sv.get("제형", ""), av["제형"]])
        for col, h in ((5, sv.get("HEX")), (7, av["HEX"])):
            if hex_ok(h):
                ws.cell(ws.max_row, col).fill = PatternFill("solid", fgColor=h.strip().lstrip("#").upper())
    for i, w in enumerate([15, 35, 30, 11, 8, 11, 8, 6, 12, 45, 20, 10, 10], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    write_aside_v8(wb, aside_rows)
    out = os.path.join(os.path.dirname(os.path.abspath(sys.argv[2])), f"compare_{datetime.now():%Y%m%d_%H%M}.xlsx")
    wb.save(out)
    print(f"비교 {len(keys)}건 → {out}")
    for v in ("거의 같음", "약간 다름", "다름", "확실히 다름", "한쪽 색 없음", "스크립트에 없음"):
        if v in stats:
            print(f"  {v}: {stats[v]}건")


if __name__ == "__main__":
    main()
