"""
올리브영 메이크업(립/베이스/아이) 색조 옵션별 대표 컬러코드 수집 PoC
- 비로그인, API key 없음, requests 계열(curl_cffi 우선) + BeautifulSoup
- 중분류(립/베이스/아이)당 상품 10개, oliveyoung.co.kr 요청 간격 5초
- 컬러칩 이미지는 메모리에서만 열어 대표색(HEX)을 뽑고 저장하지 않음
- 결과: 엑셀(.xlsx), HEX 셀에 실제 색 채움

pip install curl_cffi requests beautifulsoup4 pillow pandas openpyxl
"""
import re
import time
import random
import statistics
from io import BytesIO
from datetime import datetime
from urllib.parse import urljoin

import pandas as pd
from bs4 import BeautifulSoup
from PIL import Image
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# ── stealth: curl_cffi가 있으면 크롬 TLS 지문까지 흉내, 없으면 일반 requests ──
try:
    from curl_cffi import requests as http
    IMPERSONATE = {"impersonate": "chrome"}
except ImportError:
    import requests as http
    IMPERSONATE = {}

BASE = "https://www.oliveyoung.co.kr"
MAKEUP_URL = f"{BASE}/store/display/getCategoryShop.do?dispCatNo=10000010002"

# 중분류 이름 → dispCatNo (메이크업 페이지에서 자동으로 못 찾을 때 쓰는 기본값, DevTools로 확인 권장)
TARGETS = {
    "립메이크업": "100000100020006",
    "베이스메이크업": "100000100020001",
    "아이메이크업": "100000100020007",
}
PER_CATEGORY = 10
PAGE_DELAY = 5   # 올리브영 페이지/AJAX 요청 간격(초)
IMG_DELAY = 1    # 컬러칩 이미지(CDN) 요청 간격(초)
OUT_XLSX = "oliveyoung_makeup_colors_poc.xlsx"

# ── (선택) 브라우저 세션 이어받기 ──
# 크롬에서 올리브영 접속 → F12 → Network → 아무 문서 요청 클릭 → Request Headers의
# cookie 값과 user-agent 값을 그대로 복사해 아래에 붙여넣기 (비로그인 상태 그대로)
# → 스크립트와 같은 폴더의 cookie.txt 에 아래 형식으로 저장 (코드에 직접 넣지 않음)
#   cookie
#   <cookie 값 전체>
#   user-agent
#   <user-agent 값 전체>
import os
_HERE = os.path.dirname(os.path.abspath(__file__))
COOKIE_CANDIDATES = [os.path.join(_HERE, n) for n in ("cookie.txt", "cookie.txt.txt", "cookie")]


def _read_text(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "utf-16", "cp949"):   # BOM/메모장 저장 방식 대응
        try:
            return raw.decode(enc)
        except UnicodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


def load_browser_session(paths):
    """cookie.txt → (Cookie 헤더, UA). 같은 이름 쿠키가 여러 번 있으면 마지막(최신) 값만 사용."""
    path = next((p for p in paths if os.path.isfile(p)), None)
    if not path:
        print(f"[쿠키] 파일 없음. 찾은 폴더: {_HERE}")
        print("       폴더 내 txt 파일:", [f for f in os.listdir(_HERE) if "cookie" in f.lower()])
        return "", ""
    lines = [l.strip().lstrip("\ufeff") for l in _read_text(path).splitlines() if l.strip()]
    vals = {}
    for i, l in enumerate(lines[:-1]):
        if l.lower() in ("cookie", "user-agent"):
            vals[l.lower()] = lines[i + 1]
    jar = {}
    for part in vals.get("cookie", "").split(";"):
        if "=" in part:
            k, v = part.strip().split("=", 1)
            jar.pop(k, None)
            jar[k] = v           # 중복 시 뒤쪽(최신) 값으로 교체
    cookie = "; ".join(f"{k}={v}" for k, v in jar.items())
    print(f"[쿠키] {os.path.basename(path)} 읽음 → 쿠키 {len(jar)}개, "
          f"cf_clearance {'있음' if 'cf_clearance' in jar else '없음'}, UA {'있음' if vals.get('user-agent') else '없음'}")
    return cookie, vals.get("user-agent", "")


BROWSER_COOKIE, BROWSER_UA = load_browser_session(COOKIE_CANDIDATES)

if IMPERSONATE:
    # curl_cffi는 크롬과 같은 헤더 세트/순서를 스스로 만든다 → 최소한만 덧붙임
    HEADERS = {"Accept-Language": "ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7",
               "Referer": f"{BASE}/store/main/main.do"}
else:
    HEADERS = {
        "User-Agent": ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                       "(KHTML, like Gecko) Chrome/141.0.0.0 Safari/537.36"),
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
        "Accept-Language": "ko-KR,ko;q=0.9,en-US;q=0.8,en;q=0.7",
        "Referer": f"{BASE}/store/main/main.do",
        "Upgrade-Insecure-Requests": "1",
    }
if BROWSER_UA:
    HEADERS["User-Agent"] = BROWSER_UA      # 쿠키는 발급받은 브라우저 UA와 같아야 유효
if BROWSER_COOKIE:
    HEADERS["Cookie"] = BROWSER_COOKIE

print("HTTP 라이브러리:", "curl_cffi (chrome 위장)" if IMPERSONATE else "requests (TLS 위장 없음)")
print("브라우저 쿠키 사용:", bool(BROWSER_COOKIE))

session = http.Session()
session.headers.update(HEADERS)
_last = {"page": 0.0, "img": 0.0}


def get(url, kind="page", **kw):
    """요청 간격을 지키는 GET. 차단(403 등)이면 예외 발생."""
    delay = PAGE_DELAY if kind == "page" else IMG_DELAY
    wait = delay - (time.time() - _last[kind])
    if wait > 0:
        time.sleep(wait + random.uniform(0, 0.5))
    r = session.get(url, timeout=15, **IMPERSONATE, **kw)
    _last[kind] = time.time()
    if r.status_code == 403 or "cf-chl" in r.text[:3000]:
        server = r.headers.get("server", "?")
        raise RuntimeError(f"차단됨(status={r.status_code}, server={server}): {url}")
    r.raise_for_status()
    return r


def soup_of(resp):
    return BeautifulSoup(resp.text, "html.parser")


def first_text(soup, selector):
    el = soup.select_one(selector)
    return el.get_text(" ", strip=True) if el else ""


# ── 1. 메이크업 페이지에서 중분류 dispCatNo 찾기 ──
def find_mid_categories():
    found = {}
    try:
        soup = soup_of(get(MAKEUP_URL))
        for a in soup.select("a"):
            name = a.get_text(strip=True)
            if name not in TARGETS or name in found:
                continue
            no = a.get("data-ref-dispcatno")
            if not no:
                m = re.search(r"dispCatNo=(\d+)", a.get("href", ""))
                no = m.group(1) if m else None
            if no:
                found[name] = no
    except Exception as e:
        print(f"[WARN] 메이크업 페이지 파싱 실패, 기본값 사용: {e}")
    return {k: found.get(k, v) for k, v in TARGETS.items()}


# ── 2. 중분류 상품 목록(소분류 전체)에서 goodsNo 수집 ──
def list_goods(cat_no, limit):
    params = {"dispCatNo": cat_no, "fltDispCatNo": "", "prdSort": "01",
              "pageIdx": 1, "rowsPerPage": 24}
    soup = soup_of(get(f"{BASE}/store/display/getMCategoryList.do",
                       params=params, headers={"Referer": MAKEUP_URL}))
    goods = []
    for a in soup.select("a[data-ref-goodsno], a[href*='goodsNo=']"):
        no = a.get("data-ref-goodsno")
        if not no:
            m = re.search(r"goodsNo=(\w+)", a.get("href", ""))
            no = m.group(1) if m else None
        if no and no not in goods:
            goods.append(no)
        if len(goods) >= limit:
            break
    return goods


# ── 3. 상세 페이지 옵션(호수명 + 컬러칩 URL) 파싱 ──
OPTION_SELECTORS = ["#option_list li", "ul.sel_option_list li",
                    ".prd_option_box li", ".option_list li"]


def clean_option(text):
    text = re.sub(r"\d{1,3}(,\d{3})*\s*원", "", text)
    text = re.sub(r"\(?품절\)?|오늘드림|세일|쿠폰", "", text)
    return re.sub(r"\s+", " ", text).strip()


# 옵션명 → (호수, 호수이름)  예: "01 코랄핑크" → ("01", "코랄핑크"), "[NEW] #21N 바닐라" → ("21N", "바닐라")
HOSU_RE = re.compile(r"^#?([A-Za-z]{0,3}\d+[A-Za-z0-9.\-]*)[\s._)\]:-]*(.*)$")


def split_hosu(option_name):
    s = re.sub(r"^(\[[^\]]*\]|\([^)]*\)|\s)+", "", option_name)   # 앞쪽 [단독], (기획) 등 제거
    m = HOSU_RE.match(s)
    if m:
        return m.group(1), m.group(2).strip()
    return "", s.strip()


# sRGB(HEX) → CIELAB (D65 기준). 수집값이 아니라 HEX에서 계산하는 값
def rgb_to_lab(r, g, b):
    def lin(c):
        c /= 255
        return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
    R, G, B = lin(r), lin(g), lin(b)
    x = (0.4124564 * R + 0.3575761 * G + 0.1804375 * B) / 0.95047
    y = (0.2126729 * R + 0.7151522 * G + 0.0721750 * B) / 1.00000
    z = (0.0193339 * R + 0.1191920 * G + 0.9503041 * B) / 1.08883

    def f(t):
        return t ** (1 / 3) if t > 0.008856 else 7.787 * t + 16 / 116
    fx, fy, fz = f(x), f(y), f(z)
    return round(116 * fy - 16, 2), round(500 * (fx - fy), 2), round(200 * (fy - fz), 2)


def parse_options(soup):
    items = []
    for sel in OPTION_SELECTORS:
        items = soup.select(sel)
        if items:
            break
    out = []
    for li in items:
        label_el = li.select_one(".option_value, .txt, .tx_name") or li
        raw = re.sub(r"\s+", " ", label_el.get_text(" ", strip=True)).strip()   # 올리브영 원문 옵션명
        label = clean_option(raw)                                                 # 가격/품절 문구 제거본
        img = li.select_one("img")
        src = (img.get("data-src") or img.get("src")) if img else ""
        chip = urljoin(BASE, src) if src else ""
        if label:
            out.append((raw, label, chip))
    return out


# ── 4. 컬러칩 → 대표색(중앙부 픽셀 중앙값) ──
def chip_color(url):
    if not url:
        return "", ""
    img = Image.open(BytesIO(get(url, kind="img").content)).convert("RGB")
    w, h = img.size
    img = img.crop((int(w * .25), int(h * .25), int(w * .75), int(h * .75))).resize((16, 16))
    r, g, b = (int(statistics.median(ch)) for ch in zip(*img.getdata()))
    return f"#{r:02X}{g:02X}{b:02X}", (r, g, b)


def crawl_goods(mid_name, goods_no):
    url = f"{BASE}/store/goods/getGoodsDetail.do?goodsNo={goods_no}"
    soup = soup_of(get(url))

    name = first_text(soup, "p.prd_name, .prd_name")
    brand = first_text(soup, "#moveBrandShop, .prd_brand")
    crumbs = [a.get_text(strip=True) for a in soup.select(".page_location a.cate_y")] \
        or [a.get_text(strip=True) for a in soup.select(".page_location a")]
    crumbs = [c for c in crumbs if c and c not in ("홈", "HOME")]
    sub_cat = crumbs[-1] if crumbs else ""

    options = parse_options(soup)
    if not options:  # 옵션이 HTML에 없으면 내부 AJAX 직접 호출
        try:
            ajax = get(f"{BASE}/store/goods/getOptInfoListAjax.do",
                       params={"goodsNo": goods_no},
                       headers={"Referer": url, "X-Requested-With": "XMLHttpRequest"})
            options = parse_options(soup_of(ajax))
        except Exception as e:
            print(f"   옵션 AJAX 실패: {e}")
    if not options:
        options = [("단일/옵션없음", "단일/옵션없음", "")]

    rows = []
    for raw, label, chip in options:
        hex_code, lab = "", ("", "", "")
        try:
            if chip:
                hex_code, rgb = chip_color(chip)
                lab = rgb_to_lab(*rgb)
        except Exception as e:
            print(f"   컬러칩 실패({label}): {e}")
        hosu, hosu_name = split_hosu(label)
        rows.append({
            "상품번호": goods_no,
            "중분류": mid_name,
            "제품유형(소분류)": sub_cat,
            "브랜드명": brand,
            "상품명": name,
            "호수": hosu,
            "호수이름": hosu_name,
            "올리브영 옵션명": raw,
            "올리브영 HEX": hex_code,
            "색상(배경)": "",          # 엑셀 저장 시 HEX 색으로 칠함
            "L*": lab[0],
            "a*": lab[1],
            "b*": lab[2],
            "옵션 이미지 URL": chip,
            "수집일시": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        })
    return rows


def save_excel(rows):
    df = pd.DataFrame(rows)
    df.to_excel(OUT_XLSX, index=False)
    wb = load_workbook(OUT_XLSX)
    ws = wb.active
    hex_col = list(df.columns).index("올리브영 HEX") + 1
    fill_col = list(df.columns).index("색상(배경)") + 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, hex_col).value
        if v:
            ws.cell(r, fill_col).fill = PatternFill("solid", fgColor=v.lstrip("#"))
    widths = {"상품명": 40, "올리브영 옵션명": 30, "호수이름": 18, "옵션 이미지 URL": 50, "수집일시": 19}
    for i, c in enumerate(df.columns, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = widths.get(c, 13)
    ws.freeze_panes = "B2"
    wb.save(OUT_XLSX)
    print(f"\n저장 완료: {OUT_XLSX} ({len(df)}행)")


def main():
    try:
        get(f"{BASE}/store/main/main.do")  # 쿠키 받기용 워밍업
    except Exception as e:
        print(f"[WARN] 메인 접속 실패: {e}")

    rows = []
    for mid_name, cat_no in find_mid_categories().items():
        print(f"\n[{mid_name}] dispCatNo={cat_no}")
        try:
            goods = list_goods(cat_no, PER_CATEGORY)
        except Exception as e:
            print(f"   목록 실패: {e}")
            continue
        for i, goods_no in enumerate(goods, 1):
            try:
                got = crawl_goods(mid_name, goods_no)
                rows += got
                print(f"   {i:02d}. {goods_no} {got[0]['상품명'][:30]} → 옵션 {len(got)}개")
            except Exception as e:
                print(f"   {i:02d}. {goods_no} 실패: {e}")

    if rows:
        save_excel(rows)
    else:
        print("수집된 데이터 없음 — 차단 여부/셀렉터를 확인하세요.")


if __name__ == "__main__":
    main()
